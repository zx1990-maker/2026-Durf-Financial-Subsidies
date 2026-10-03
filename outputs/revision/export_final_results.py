#!/usr/bin/env python3
"""Export final core results with AMOUNT gap (gap × total assets)."""
import sys, json, warnings, numpy as np, pandas as pd
from pathlib import Path
import openpyxl

warnings.filterwarnings('ignore')
PROJECT_DIR = Path(__file__).resolve().parent.parent.parent
OUTPUT_DIR = PROJECT_DIR / "final_results"
OUTPUT_DIR.mkdir(exist_ok=True)
INDEX_DIR = PROJECT_DIR / "input_index_weight"

sys.path.insert(0, str(PROJECT_DIR))
from run_counterfactual import *
OUTPUT_DIR = PROJECT_DIR / "final_results"  # 重新断言：import * 会把 OUTPUT_DIR 覆盖为 outputs/
from sklearn.ensemble import GradientBoostingRegressor
from sklearn.impute import SimpleImputer
from sklearn.preprocessing import StandardScaler, OneHotEncoder

RANDOM_STATE=42; SOE={'中央国有企业','地方国有企业'}

# ============================================================================
# 1. Load + screen (same as final: ST / insolvent / extreme excluded)
# ============================================================================
print('Loading...')
all_a = load_all_a_shares(); all_a = clean_and_filter(all_a, 'A')
st_firms = set(all_a[all_a['firm_name'].str.contains('ST', na=False)]['firm_id'].unique())
all_a = all_a[~all_a['firm_id'].isin(st_firms)]
all_a = all_a[~(all_a['FinancialLeverage'] > 1.0)]
all_a = all_a[all_a['EBI_A'].abs() <= 0.5]
all_a['is_soe'] = all_a['ownership'].isin(SOE)
all_a['is_private'] = all_a['ownership']=='民营企业'
train_pool = all_a[all_a['is_private']]
print(f'Train pool: {len(train_pool)} obs, {train_pool["firm_id"].nunique()} firms')

# Index codes
idx_files = {
    'CSI300': ('沪深300-成分及权重-20260811.xlsx', '沪深300-成分进出记录-20260811.xlsx'),
    'CSI500': ('中证500-成分及权重-20260811.xlsx', '中证500-成分进出记录-20260811.xlsx'),
    'CSI1000': ('中证1000-成分及权重-20260811.xlsx', '中证1000-成分进出记录-20260811.xlsx'),
}
codes={}; all_records={}
for idx,(fw,fe) in idx_files.items():
    wb=openpyxl.load_workbook(INDEX_DIR/fw, read_only=True, data_only=True); ws=wb.active
    codes[idx]={str(r[0]).strip() for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and str(r[2]).strip()!='金融'}
    wb.close()
    wb2=openpyxl.load_workbook(INDEX_DIR/fe, read_only=True, data_only=True); ws2=wb2.active
    all_records[idx]=sorted([{'date':str(r[0])[:10],'code':str(r[1]).strip(),'action':str(r[7]).strip()}
                             for r in ws2.iter_rows(min_row=2, values_only=True) if r[0] and r[1] and r[7]],
                            key=lambda x:x['date'])
    wb2.close()
yearly_codes={}
for idx in ['CSI300','CSI500','CSI1000']:
    cur=codes[idx]; rev=sorted(all_records[idx],key=lambda r:r['date'],reverse=True)
    yearly_codes[idx]={}
    for y in range(2019,2026):
        td=f'{y}-07-01'; cs=set(cur)
        for rec in rev:
            if rec['date']<=td: break
            if rec['action']=='纳入': cs.discard(rec['code'])
            elif rec['action']=='剔除': cs.add(rec['code'])
        yearly_codes[idx][y]=cs

feats=[f for f in NUMERIC_FEATURES if f in all_a.columns]

# ============================================================================
# 2. Rolling: GBM predict, store firm-level gap + AMOUNT gap
# ============================================================================
firm_rows = []

for fc_year in range(2019, 2026):
    train = train_pool[train_pool['target_year']<fc_year].copy()  # 严格 ex-ante：结果年份 < 预测年
    tv = train.dropna(subset=feats+['EBI_A']).copy()
    for c in feats: lo,hi=tv[c].quantile(0.05),tv[c].quantile(0.95); tv[c+'_w']=tv[c].clip(lo,hi)
    yl,yh=tv['EBI_A'].quantile(0.05),tv['EBI_A'].quantile(0.95); tv['EBI_A_w']=tv['EBI_A'].clip(yl,yh)
    ncw=[c+'_w' for c in feats]
    imp=SimpleImputer(strategy='median'); scl=StandardScaler()
    Xt=scl.fit_transform(imp.fit_transform(tv[ncw].values))
    ohe=OneHotEncoder(handle_unknown='ignore',sparse_output=False)
    Xt=np.hstack([Xt,ohe.fit_transform(tv[['industry_l2']].fillna('Unknown'))])
    yt=tv['EBI_A_w'].values

    m = GradientBoostingRegressor(n_estimators=200,max_depth=4,learning_rate=0.05,min_samples_leaf=10,random_state=RANDOM_STATE)
    m.fit(Xt, yt)

    soe_all = all_a[(all_a['is_soe'])&(all_a['target_year']==fc_year)]
    targets = {'All A-share SOE': soe_all}
    for idx in ['CSI300','CSI500','CSI1000']:
        targets[f'{idx} SOE'] = soe_all[soe_all['firm_id'].isin(yearly_codes[idx][fc_year])]

    for tn, pdf in targets.items():
        pv = pdf.copy()
        for c in feats: lo,hi=tv[c].quantile(0.05),tv[c].quantile(0.95); pv[c+'_w']=pv[c].clip(lo,hi)
        pd2 = pv.dropna(subset=ncw)
        if len(pd2)==0: continue
        Xp = np.hstack([scl.transform(imp.transform(pd2[ncw].values)),
                        ohe.transform(pd2[['industry_l2']].fillna('Unknown'))])
        preds = m.predict(Xp)
        actual_ebia = pd2['EBI_A'].values
        assets_target = pd2['资产总计_target'].values  # t+1 total assets (元)
        gap_ratio = preds - actual_ebia
        # AMOUNT gap = gap × total assets (元)
        amount_gap_yuan = gap_ratio * assets_target

        for i in range(len(pd2)):
            firm_rows.append({
                'year': fc_year,
                'sample': tn,
                'firm_id': pd2['firm_id'].iloc[i],
                'firm_name': pd2['firm_name'].iloc[i] if 'firm_name' in pd2.columns else '',
                'industry_l2': pd2['industry_l2'].iloc[i],
                'ownership': pd2['ownership'].iloc[i],
                'total_assets_target_yuan': assets_target[i],
                'actual_ebia': actual_ebia[i],
                'predicted_ebia': preds[i],
                'gap_ratio': gap_ratio[i],
                'gap_amount_yuan': amount_gap_yuan[i],
                'gap_amount_yi': amount_gap_yuan[i] / 1e8,  # 亿元
            })
    print(f'  {fc_year} done')

firm_df = pd.DataFrame(firm_rows)
firm_df.to_csv(OUTPUT_DIR/'firm_level_gap_all_years.csv', index=False, encoding='utf-8-sig')
print(f'\n✓ firm_level_gap_all_years.csv ({len(firm_df)} rows)')

# ============================================================================
# 3. 2025 detailed + summaries
# ============================================================================
firm_2025 = firm_df[firm_df['year']==2025].copy()
firm_2025.to_csv(OUTPUT_DIR/'firm_level_gap_2025.csv', index=False, encoding='utf-8-sig')
print(f'✓ firm_level_gap_2025.csv ({len(firm_2025)} rows)')

# Summary by sample (2025)
sum_sample = firm_2025.groupby('sample').agg(
    n_firms=('firm_id','count'),
    total_assets_yi=('total_assets_target_yuan', lambda x: x.sum()/1e8),
    mean_gap_ratio=('gap_ratio','mean'),
    sum_gap_amount_yi=('gap_amount_yi','sum'),
    mean_gap_amount_yi=('gap_amount_yi','mean'),
    median_gap_amount_yi=('gap_amount_yi','median'),
).reset_index()
sum_sample.to_csv(OUTPUT_DIR/'summary_by_sample_2025.csv', index=False, encoding='utf-8-sig')
print(f'✓ summary_by_sample_2025.csv')

# Summary by industry (2025, All A-share SOE)
all_2025 = firm_2025[firm_2025['sample']=='All A-share SOE']
sum_ind = all_2025.groupby('industry_l2').agg(
    n_firms=('firm_id','count'),
    mean_gap_ratio=('gap_ratio','mean'),
    median_gap_ratio=('gap_ratio','median'),
    sum_gap_amount_yi=('gap_amount_yi','sum'),
    mean_gap_amount_yi=('gap_amount_yi','mean'),
).reset_index().sort_values('sum_gap_amount_yi', ascending=False)
sum_ind.to_csv(OUTPUT_DIR/'summary_by_industry_2025.csv', index=False, encoding='utf-8-sig')
print(f'✓ summary_by_industry_2025.csv')

# Summary by year (All A-share SOE)
all_years = firm_df[firm_df['sample']=='All A-share SOE']
sum_year = all_years.groupby('year').agg(
    n_firms=('firm_id','count'),
    mean_gap_ratio=('gap_ratio','mean'),
    sum_gap_amount_yi=('gap_amount_yi','sum'),
    mean_gap_amount_yi=('gap_amount_yi','mean'),
).reset_index()
sum_year.to_csv(OUTPUT_DIR/'summary_by_year_all_soe.csv', index=False, encoding='utf-8-sig')
print(f'✓ summary_by_year_all_soe.csv')

# ============================================================================
# 4. Print key results
# ============================================================================
print('\n' + '='*90)
print('  核心结果：金额 Gap（gap × 资产总计，亿元）')
print('='*90)

print('\n【分样本 2025】')
print(sum_sample.to_string(index=False))

print('\n【分年份 All A-share SOE】')
print(sum_year.to_string(index=False))

print('\n【分行业 2025（All SOE，按金额 gap 降序）】')
print(sum_ind.head(15).to_string(index=False))

print('\n✓ 全部核心结果已导出到 final_results/')
