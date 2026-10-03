#!/usr/bin/env python3
"""Rerun full analysis with screened sample (exclude ST/insolvent/extreme)."""
import sys, json, warnings, numpy as np, pandas as pd
from pathlib import Path
import openpyxl

warnings.filterwarnings('ignore')
PROJECT_DIR = Path(__file__).resolve().parent.parent.parent
OUTPUT_DIR = Path(__file__).resolve().parent
INDEX_DIR = PROJECT_DIR / "input_index_weight"

sys.path.insert(0, str(PROJECT_DIR))
from run_counterfactual import *
OUTPUT_DIR = Path(__file__).resolve().parent  # 重新断言：import * 会把 OUTPUT_DIR 覆盖为 outputs/
from sklearn.linear_model import Ridge
from sklearn.ensemble import RandomForestRegressor, GradientBoostingRegressor
from sklearn.impute import SimpleImputer
from sklearn.preprocessing import StandardScaler, OneHotEncoder
from sklearn.model_selection import GroupKFold, cross_val_predict
from sklearn.metrics import r2_score, mean_squared_error

RANDOM_STATE=42; RIDGE_ALPHA=10.0; SOE={'中央国有企业','地方国有企业'}

# ============================================================================
# 1. Load + screen
# ============================================================================
print('Loading...')
all_a = load_all_a_shares(); all_a = clean_and_filter(all_a, 'A')

# --- Screen 1: ST firms (firm-level, current stock_name) ---
n_before = all_a['firm_id'].nunique()
st_firms = set(all_a[all_a['firm_name'].str.contains('ST', na=False)]['firm_id'].unique())
all_a = all_a[~all_a['firm_id'].isin(st_firms)]
print(f'Screen 1 ST: removed {len(st_firms)} firms')

# --- Screen 2: insolvent (FinancialLeverage > 1, firm-year) ---
# 只删除 FinancialLeverage > 1 的行，保留 NaN（数据缺失 ≠ 资不抵债）
ins_before = len(all_a)
all_a = all_a[~(all_a['FinancialLeverage'] > 1.0)]
print(f'Screen 2 insolvent: removed {ins_before - len(all_a)} obs')

# --- Screen 3: extreme EBI/A (|EBI/A| > 0.5, firm-year) ---
ext_before = len(all_a)
all_a = all_a[all_a['EBI_A'].abs() <= 0.5]
print(f'Screen 3 extreme EBI/A: removed {ext_before - len(all_a)} obs')

n_after = all_a['firm_id'].nunique()
print(f'Firms: {n_before} → {n_after} ({n_after/n_before*100:.1f}% retained)')

all_a['is_soe'] = all_a['ownership'].isin(SOE)
all_a['is_private'] = all_a['ownership'] == '民营企业'
train_pool = all_a[all_a['is_private']]
print(f'Train pool: {len(train_pool)} obs, {train_pool["firm_id"].nunique()} firms')
print(f'SOE pool: {all_a["is_soe"].sum()} obs, {all_a[all_a["is_soe"]]["firm_id"].nunique()} firms')

# ============================================================================
# 2. Index codes
# ============================================================================
idx_files = {
    'CSI300': ('沪深300-成分及权重-20260811.xlsx', '沪深300-成分进出记录-20260811.xlsx'),
    'CSI500': ('中证500-成分及权重-20260811.xlsx', '中证500-成分进出记录-20260811.xlsx'),
    'CSI1000': ('中证1000-成分及权重-20260811.xlsx', '中证1000-成分进出记录-20260811.xlsx'),
}
codes={}; all_records={}
for idx,(fw,fe) in idx_files.items():
    wb=openpyxl.load_workbook(INDEX_DIR/fw, read_only=True, data_only=True); ws=wb.active
    codes[idx]={str(r[0]).strip() for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and str(r[2]).strip()!='金融'}
    wb.close()
    wb2=openpyxl.load_workbook(INDEX_DIR/fe, read_only=True, data_only=True); ws2=wb2.active
    all_records[idx]=sorted([{'date':str(r[0])[:10],'code':str(r[1]).strip(),'action':str(r[7]).strip()}
                             for r in ws2.iter_rows(min_row=2, values_only=True) if r[0] and r[1] and r[7]],
                            key=lambda x:x['date'])
    wb2.close()
yearly_codes={}
for idx in ['CSI300','CSI500','CSI1000']:
    cur=codes[idx]; rev=sorted(all_records[idx],key=lambda r:r['date'],reverse=True)
    yearly_codes[idx]={}
    for y in range(2019,2026):
        td=f'{y}-07-01'; cs=set(cur)
        for rec in rev:
            if rec['date']<=td: break
            if rec['action']=='纳入': cs.discard(rec['code'])
            elif rec['action']=='剔除': cs.add(rec['code'])
        yearly_codes[idx][y]=cs

# ============================================================================
# 3. Models
# ============================================================================
def build(mt):
    if mt=='ridge': return Ridge(alpha=RIDGE_ALPHA, random_state=RANDOM_STATE)
    elif mt=='random_forest': return RandomForestRegressor(n_estimators=300,max_depth=8,min_samples_leaf=10,max_features='sqrt',random_state=RANDOM_STATE,n_jobs=-1)
    else: return GradientBoostingRegressor(n_estimators=200,max_depth=4,learning_rate=0.05,min_samples_leaf=10,random_state=RANDOM_STATE)

def cse(gaps,fids):
    u=np.unique(fids); fm=np.array([gaps[fids==f].mean() for f in u])
    return np.std(fm,ddof=1)/np.sqrt(len(u)) if len(u)>1 else np.nan, len(u)

feats=[f for f in NUMERIC_FEATURES if f in all_a.columns]
results=[]

for fc_year in range(2019,2026):
    train=train_pool[train_pool['target_year']<fc_year].copy()  # 严格 ex-ante：结果年份 < 预测年
    train_yrs=sorted(train['target_year'].unique())
    tv=train.dropna(subset=feats+['EBI_A']).copy()
    for c in feats: lo,hi=tv[c].quantile(0.05),tv[c].quantile(0.95); tv[c+'_w']=tv[c].clip(lo,hi)
    yl,yh=tv['EBI_A'].quantile(0.05),tv['EBI_A'].quantile(0.95); tv['EBI_A_w']=tv['EBI_A'].clip(yl,yh)
    ncw=[c+'_w' for c in feats]
    imp=SimpleImputer(strategy='median'); scl=StandardScaler()
    Xt=scl.fit_transform(imp.fit_transform(tv[ncw].values))
    ohe=OneHotEncoder(handle_unknown='ignore',sparse_output=False)
    Xt=np.hstack([Xt,ohe.fit_transform(tv[['industry_l2']].fillna('Unknown'))])
    yt=tv['EBI_A_w'].values; grp=tv['firm_id'].values

    soe_all=all_a[(all_a['is_soe'])&(all_a['target_year']==fc_year)]
    targets={'All A-share SOE':soe_all}
    for idx in ['CSI300','CSI500','CSI1000']:
        targets[f'{idx} SOE']=soe_all[soe_all['firm_id'].isin(yearly_codes[idx][fc_year])]

    for mt in ['ridge','random_forest','gbm']:
        nf=tv['firm_id'].nunique(); ns=min(5,nf)
        if ns>=2:
            oof_p=cross_val_predict(build(mt),Xt,yt,cv=GroupKFold(ns).split(Xt,yt,groups=grp),n_jobs=-1 if mt=='random_forest' else 1)
            oof_r2=r2_score(yt,oof_p); oof_rmse=np.sqrt(mean_squared_error(yt,oof_p))
        else: oof_r2=oof_rmse=np.nan
        m=build(mt); m.fit(Xt,yt)
        for tn,pdf in targets.items():
            if len(pdf)<3: continue
            pv=pdf.copy()
            for c in feats: lo,hi=tv[c].quantile(0.05),tv[c].quantile(0.95); pv[c+'_w']=pv[c].clip(lo,hi)
            pd2=pv.dropna(subset=ncw)
            if len(pd2)==0: continue
            Xp=np.hstack([scl.transform(imp.transform(pd2[ncw].values)),ohe.transform(pd2[['industry_l2']].fillna('Unknown'))])
            preds=m.predict(Xp); actuals=pd2['EBI_A'].values; gaps=preds-actuals
            se,nc=cse(gaps,pd2['firm_id'].values); mg=np.mean(gaps)
            results.append({'forecast_year':fc_year,'model':mt,'target':tn,
                'n_train_obs':len(tv),'n_train_firms':nf,'n_pred':len(pd2),
                'pred_mean':np.mean(preds),'actual_mean':np.mean(actuals),'mean_gap':mg,
                'clustered_se':se,'ci_95_lower':mg-1.96*se if not np.isnan(se) else np.nan,
                'ci_95_upper':mg+1.96*se if not np.isnan(se) else np.nan,
                'oof_r2':oof_r2,'oof_rmse':oof_rmse})
    print(f'  {fc_year}: train={train_yrs[0]}–{train_yrs[-1]} ({len(tv)} obs), All SOE={len(soe_all)}, 300={len(targets.get("CSI300 SOE",pd.DataFrame()))}, 500={len(targets.get("CSI500 SOE",pd.DataFrame()))}, 1000={len(targets.get("CSI1000 SOE",pd.DataFrame()))}')

df=pd.DataFrame(results)
df.to_csv(OUTPUT_DIR/'filtered_rolling.csv', index=False, encoding='utf-8-sig')
print('\n✓ filtered_rolling.csv saved')

# Print GBM summary
print('\n' + '='*85)
print('  GBM (main) — Gap by Target × Year (SCREENED sample)')
print('='*85)
gbm=df[df['model']=='gbm']
print('{:<6} {:<16} {:<16} {:<16} {:<16}'.format('Year','All SOE','CSI300','CSI500','CSI1000'))
print('-'*70)
for year in range(2019,2026):
    row='{:<6}'.format(year)
    for tn in ['All A-share SOE','CSI300 SOE','CSI500 SOE','CSI1000 SOE']:
        r=gbm[(gbm['forecast_year']==year)&(gbm['target']==tn)]
        if len(r)>0: row+=' {:>+8.4f} ({:>3d})'.format(r.iloc[0]['mean_gap'],int(r.iloc[0]['n_pred']))
        else: row+='        -    '
    print(row)

# Ridge summary
print('\n  RIDGE — Gap by Target × Year (SCREENED)')
print('='*85)
ridge=df[df['model']=='ridge']
print('{:<6} {:<16} {:<16} {:<16} {:<16}'.format('Year','All SOE','CSI300','CSI500','CSI1000'))
print('-'*70)
for year in range(2019,2026):
    row='{:<6}'.format(year)
    for tn in ['All A-share SOE','CSI300 SOE','CSI500 SOE','CSI1000 SOE']:
        r=ridge[(ridge['forecast_year']==year)&(ridge['target']==tn)]
        if len(r)>0: row+=' {:>+8.4f} ({:>3d})'.format(r.iloc[0]['mean_gap'],int(r.iloc[0]['n_pred']))
        else: row+='        -    '
    print(row)

# RF summary
print('\n  RF — Gap by Target × Year (SCREENED)')
print('='*85)
rf=df[df['model']=='random_forest']
print('{:<6} {:<16} {:<16} {:<16} {:<16}'.format('Year','All SOE','CSI300','CSI500','CSI1000'))
print('-'*70)
for year in range(2019,2026):
    row='{:<6}'.format(year)
    for tn in ['All A-share SOE','CSI300 SOE','CSI500 SOE','CSI1000 SOE']:
        r=rf[(rf['forecast_year']==year)&(rf['target']==tn)]
        if len(r)>0: row+=' {:>+8.4f} ({:>3d})'.format(r.iloc[0]['mean_gap'],int(r.iloc[0]['n_pred']))
        else: row+='        -    '
    print(row)

print('\n✓ Screening rerun complete')
